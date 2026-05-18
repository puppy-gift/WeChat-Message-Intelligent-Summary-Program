#!/usr/bin/env python3
"""
微信聊天记录 AI 智能总结工具 - 三合一版本
模式1: 自动抓取 (wx-cli)
模式2: 文件上传 (解析多种格式)
模式3: 手动粘贴 (粘贴聊天记录)
"""

import sys, os, subprocess, shutil, json, datetime, webbrowser, argparse, re, tempfile
from flask import Flask, render_template, request, jsonify, Response
from openai import OpenAI
from cryptography.fernet import Fernet

# ==================== 全局错误捕获（确保闪退时记录日志） ====================
def log_crash(exc):
    try:
        log_dir = os.path.dirname(os.path.abspath(__file__)) if not getattr(sys, 'frozen', False) else os.path.dirname(sys.executable)
        with open(os.path.join(log_dir, "crash.log"), "a", encoding="utf-8") as f:
            f.write(f"{datetime.datetime.now()} - CRASH: {exc}\n")
    except:
        pass

try:
    # ==================== 管理员权限提升 ====================
    def is_admin():
        try:
            import ctypes
            return ctypes.windll.shell32.IsUserAnAdmin()
        except:
            return False

    if sys.platform == "win32" and not is_admin():
        import ctypes
        params = " ".join(sys.argv[1:]) if len(sys.argv) > 1 else ""
        ctypes.windll.shell32.ShellExecuteW(None, "runas", sys.executable, params, None, 1)
        sys.exit(0)

    app = Flask(__name__)

    # ==================== 动态路径 ====================
    if getattr(sys, 'frozen', False):
        BASE_DIR = os.path.dirname(sys.executable)
    else:
        BASE_DIR = os.path.dirname(os.path.abspath(__file__))

    PROVIDERS = {
        "deepseek": {"base_url": "https://api.deepseek.com", "default_model": "deepseek-chat"},
        "openai": {"base_url": "https://api.openai.com/v1", "default_model": "gpt-4o"},
        "grok": {"base_url": "https://api.x.ai/v1", "default_model": "grok-2-latest"},
        "gemini": {"base_url": "https://generativelanguage.googleapis.com/v1beta/openai/", "default_model": "gemini-2.0-flash"},
    }

    MAX_CHUNK_TOKENS = 3000
    FINAL_SUMMARY_TOKENS = 4000

    # wx-cli 相关（仅模式1）
    WX_CLI_PATH = os.path.join(BASE_DIR, "wx-cli.exe")
    WX_CONFIG_FILE = os.path.join(BASE_DIR, "config.json")
    WX_ALL_KEYS_FILE = os.path.join(BASE_DIR, "all_keys.json")
    WX_CONFIG_DIR = os.path.join(os.path.expanduser("~"), ".wx-cli")

    # 加密存储
    KEY_FILE = os.path.join(BASE_DIR, "secret.key")
    CONFIG_FILE = os.path.join(BASE_DIR, "config.enc")

    # ==================== 加密工具 ====================
    def generate_key():
        key = Fernet.generate_key()
        with open(KEY_FILE, "wb") as f:
            f.write(key)

    def load_fernet():
        if not os.path.exists(KEY_FILE):
            generate_key()
        return Fernet(open(KEY_FILE, "rb").read())

    def load_api_keys():
        if not os.path.exists(CONFIG_FILE):
            return {}
        fernet = load_fernet()
        with open(CONFIG_FILE, "rb") as f:
            encrypted = f.read()
        try:
            return json.loads(fernet.decrypt(encrypted).decode())
        except:
            return {}

    def save_api_keys(keys_dict):
        fernet = load_fernet()
        encrypted = fernet.encrypt(json.dumps(keys_dict).encode())
        with open(CONFIG_FILE, "wb") as f:
            f.write(encrypted)

    def get_api_key(provider):
        return load_api_keys().get(provider, "")

    def set_api_key(provider, key):
        keys = load_api_keys()
        keys[provider] = key
        save_api_keys(keys)

    # ==================== wx-cli 释放与初始化（仅模式1） ====================
    def ensure_wx_cli():
        """打包后自动释放 wx-cli.exe 到 EXE 同目录"""
        if not getattr(sys, 'frozen', False):
            return
        meipass = sys._MEIPASS
        src = os.path.join(meipass, "wx-cli.exe")
        dst = WX_CLI_PATH
        if not os.path.exists(dst) and os.path.exists(src):
            shutil.copy2(src, dst)

    def auto_init_wx_cli():
        """自动运行 wx-cli init 生成密钥文件"""
        if os.path.exists(WX_ALL_KEYS_FILE):
            return True
        print("未检测到微信密钥，正在自动初始化 wx-cli ...")
        try:
            result = subprocess.run(
                [WX_CLI_PATH, "init"], capture_output=True, text=True, encoding='utf-8',
                cwd=BASE_DIR, timeout=120
            )
            print(result.stdout)
            if os.path.exists(WX_ALL_KEYS_FILE):
                print("密钥文件生成成功")
                return True
        except Exception as e:
            print(f"自动初始化失败: {e}")
        return False

    def init_wx_config():
        """确保程序目录有微信配置文件，并预热 daemon"""
        ensure_wx_cli()
        if not os.path.exists(WX_CLI_PATH):
            print("提示：未找到 wx-cli.exe，自动抓取功能不可用。文件上传和手动粘贴仍可正常使用。")
            return

        # 尝试自动生成密钥，如果失败则从用户目录复制
        if not auto_init_wx_cli():
            src_keys = os.path.join(WX_CONFIG_DIR, "all_keys.json")
            if os.path.exists(src_keys):
                shutil.copy2(src_keys, WX_ALL_KEYS_FILE)
                print("已从用户目录复制 all_keys.json")
            else:
                print("警告：无法获取微信密钥文件，自动抓取功能可能无法使用")
                return

        # 复制 config.json
        src_config = os.path.join(WX_CONFIG_DIR, "config.json")
        if not os.path.exists(WX_CONFIG_FILE) and os.path.exists(src_config):
            shutil.copy2(src_config, WX_CONFIG_FILE)
            print("已复制 config.json")

        # 预热 daemon
        print("正在启动微信 daemon...")
        try:
            subprocess.run(
                [WX_CLI_PATH, "sessions"], capture_output=True, text=True, encoding='utf-8',
                cwd=BASE_DIR, timeout=120
            )
            print("微信 daemon 启动成功")
        except subprocess.TimeoutExpired:
            print("警告：daemon 启动超时，但可继续使用")
        except Exception as e:
            print(f"警告：daemon 预热失败：{e}")

    def run_wx_cli(args, timeout=60):
        """调用 wx-cli 并返回输出"""
        try:
            result = subprocess.run(
                [WX_CLI_PATH] + args,
                capture_output=True, text=True, encoding='utf-8',
                timeout=timeout, cwd=BASE_DIR, check=True
            )
            return result.stdout
        except subprocess.CalledProcessError as e:
            raise RuntimeError(e.stderr.strip() or "wx-cli 调用失败")
        except FileNotFoundError:
            raise RuntimeError("找不到 wx-cli.exe")

    # ==================== 多格式文件解析 ====================
    def extract_text_from_docx(file_stream):
        import docx
        doc = docx.Document(file_stream)
        return "\n".join([para.text for para in doc.paragraphs if para.text.strip()])

    def extract_text_from_doc(file_stream):
        import doc2txt
        with tempfile.NamedTemporaryFile(delete=False, suffix=".doc") as tmp:
            tmp.write(file_stream.read())
            tmp_path = tmp.name
        try:
            text = doc2txt.process(tmp_path)
            return text or ""
        finally:
            os.unlink(tmp_path)

    def extract_text_from_pdf(file_stream):
        import pdfplumber
        text_parts = []
        with pdfplumber.open(file_stream) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def extract_text_from_xlsx(file_stream):
        import openpyxl
        wb = openpyxl.load_workbook(file_stream, read_only=True)
        rows = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                if row_text.strip():
                    rows.append(row_text)
        wb.close()
        return "\n".join(rows)

    def extract_text_from_xls(file_stream):
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_stream.read())
        rows = []
        for sheet in wb.sheets():
            for row_idx in range(sheet.nrows):
                row_text = "\t".join([str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)])
                if row_text.strip():
                    rows.append(row_text)
        return "\n".join(rows)

    def extract_text_from_odt(file_stream):
        from odfdo import Document as OdfDocument
        with tempfile.NamedTemporaryFile(delete=False, suffix=".odt") as tmp:
            tmp.write(file_stream.read())
            tmp_path = tmp.name
        try:
            doc = OdfDocument(tmp_path)
            return doc.get_formatted_text()
        finally:
            os.unlink(tmp_path)

    def parse_chat_file(file):
        """根据文件扩展名选择解析方式，提取文本并解析聊天记录格式"""
        filename = file.filename.lower()
        file.seek(0)

        # 第1步：根据扩展名提取纯文本
        if filename.endswith(".docx"):
            raw_text = extract_text_from_docx(file)
        elif filename.endswith(".doc"):
            raw_text = extract_text_from_doc(file)
        elif filename.endswith(".pdf"):
            raw_text = extract_text_from_pdf(file)
        elif filename.endswith(".xlsx"):
            raw_text = extract_text_from_xlsx(file)
        elif filename.endswith(".xls"):
            raw_text = extract_text_from_xls(file)
        elif filename.endswith(".csv"):
            raw_text = file.read().decode("utf-8", errors="ignore")
        elif filename.endswith(".odt") or filename.endswith(".ods"):
            raw_text = extract_text_from_odt(file)
        elif filename.endswith(".rtf"):
            content = file.read().decode("utf-8", errors="ignore")
            raw_text = re.sub(r'\\[a-z]+\d*', '', content)
            raw_text = re.sub(r'[{}]', '', raw_text)
        elif filename.endswith((".html", ".htm")):
            content = file.read().decode("utf-8", errors="ignore")
            raw_text = re.sub(r'<[^>]+>', '', content)
            raw_text = re.sub(r'&[a-z]+;', ' ', raw_text)
        elif filename.endswith((".md", ".markdown")):
            raw_text = file.read().decode("utf-8", errors="ignore")
        elif filename.endswith(".json"):
            content = file.read().decode("utf-8", errors="ignore")
            try:
                data = json.loads(content)
                raw_text = json.dumps(data, ensure_ascii=False, indent=2)
            except:
                raw_text = content
        elif filename.endswith(".xml"):
            content = file.read().decode("utf-8", errors="ignore")
            raw_text = re.sub(r'<[^>]+>', '', content)
        elif filename.endswith(".txt"):
            try:
                raw_text = file.read().decode("utf-8")
            except UnicodeDecodeError:
                file.seek(0)
                raw_text = file.read().decode("gbk", errors="ignore")
        else:
            # 其他格式尝试作为纯文本
            try:
                raw_text = file.read().decode("utf-8")
            except UnicodeDecodeError:
                file.seek(0)
                raw_text = file.read().decode("gbk", errors="ignore")

        # 第2步：解析聊天记录格式
        messages = []
        for line in raw_text.splitlines():
            line = line.strip()
            if not line:
                continue
            match = re.match(r"\[?(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})\]? (\S+): (.+)", line)
            if match:
                time_str, sender, msg_content = match.groups()
                messages.append({
                    "sender": sender,
                    "content": msg_content,
                    "time": _parse_timestamp(time_str)
                })
            else:
                messages.append({
                    "sender": "未知",
                    "content": line,
                    "time": 0
                })
        return messages

    def _parse_timestamp(time_str):
        try:
            dt = datetime.datetime.strptime(time_str.strip(), "%Y-%m-%d %H:%M:%S")
            return int(dt.timestamp())
        except:
            return 0

    # ==================== AI 总结核心 ====================
    def estimate_token_count(text):
        return len(text) // 3

    def split_text_into_chunks(text, max_tokens):
        sentences = text.replace("\n", " ").split("。")
        chunks = []
        current = ""
        for sent in sentences:
            sent = sent.strip()
            if not sent:
                continue
            tentative = current + "。" + sent if current else sent
            if estimate_token_count(tentative) <= max_tokens:
                current = tentative
            else:
                if current:
                    chunks.append(current)
                current = sent
        if current:
            chunks.append(current)
        return chunks

    def create_client(provider, api_key=None):
        config = PROVIDERS.get(provider)
        if not config:
            raise ValueError(f"不支持的提供商: {provider}")
        key = api_key or get_api_key(provider)
        if not key:
            raise RuntimeError(f"请提供 {provider} 的 API 密钥")
        return OpenAI(api_key=key, base_url=config["base_url"]), config["default_model"]

    def build_system_prompt(style, keywords=None):
        base = (
            "你是一个专业的聊天记录总结助手。请根据提供的聊天记录，生成一个条理清晰的总结，"
            "必须包含以下内容：\n"
            "1. 与我（提问者）直接相关的讨论或提到我的事情；\n"
            "2. 聊天中出现的所有关键信息（决定、计划、重要通知等）；\n"
            "3. 如果聊天中有多人约定之后要做某事，请明确列出谁、什么时候、做什么。\n"
            "总结应简明扼要，避免废话。"
        )
        if keywords:
            kw_str = "、".join(keywords)
            base += f"\n\n特别注意：请重点总结并突出聊天记录中提及“{kw_str}”的相关消息。"
        if not style:
            style_instruction = "请根据聊天内容的语气和氛围，自行选择合适的总结风格。"
        else:
            style_instruction = f"总结的整体风格请严格遵循以下描述：{style}"
        return base + "\n" + style_instruction

    def summarize_chunk(client, model, chunk, style, keywords=None):
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": build_system_prompt(style, keywords)},
                {"role": "user", "content": f"请总结以下聊天记录片段：\n\n{chunk}"},
            ],
            temperature=0.3,
        )
        content = response.choices[0].message.content.strip()
        return content or "AI 未能从这段记录中提取到有效总结。"

    def merge_summaries(client, model, summaries, style, keywords=None):
        combined = "\n\n---\n\n".join(summaries)
        if estimate_token_count(combined) > FINAL_SUMMARY_TOKENS:
            sub_chunks = split_text_into_chunks(combined, FINAL_SUMMARY_TOKENS)
            sub_summaries = [summarize_chunk(client, model, c, style, keywords) for c in sub_chunks]
            combined = "\n\n---\n\n".join(sub_summaries)
        response = client.chat.completions.create(
            model=model,
            messages=[
                {"role": "system", "content": build_system_prompt(style, keywords)},
                {"role": "user", "content": f"以下是多段聊天记录的摘要，请整合成一份最终总结：\n\n{combined}"},
            ],
            temperature=0.3,
        )
        return response.choices[0].message.content.strip()

    def sse_data(text):
        for line in text.split('\n'):
            yield f"data: {line}\n"
        yield "\n"

    # ==================== 路由 ====================
    @app.route("/")
    def index():
        return render_template("index.html", providers=list(PROVIDERS.keys()))

    @app.route("/wx/status")
    def wx_status():
        try:
            run_wx_cli(["sessions"], timeout=20)
            return jsonify({"online": True})
        except Exception as e:
            return jsonify({"online": False, "error": str(e)})

    @app.route("/wx/chats")
    def list_chats():
        try:
            import yaml
            out = run_wx_cli(["sessions"], timeout=30)
            data = yaml.safe_load(out)
            chats = []
            if isinstance(data, list):
                for item in data:
                    chats.append({
                        "id": item.get("username", ""),
                        "name": item.get("chat", "未知"),
                        "lastMsg": item.get("summary", ""),
                        "time": item.get("timestamp", 0)
                    })
            return jsonify(chats)
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/wx/messages/<chat_id>")
    def get_messages(chat_id):
        limit = request.args.get("limit", 500)
        since = request.args.get("since", "")
        until = request.args.get("until", "")
        args = ["history", chat_id, "--limit", str(limit), "--json"]
        if since:
            args += ["--since", since]
        if until:
            args += ["--until", until]
        try:
            out = run_wx_cli(args, timeout=30)
            data = json.loads(out)
            msg_list = data if isinstance(data, list) else data.get("messages", [])
            messages = []
            for msg in msg_list:
                messages.append({
                    "sender": msg.get("sender", "未知"),
                    "content": msg.get("content", ""),
                    "time": msg.get("createTime", 0)
                })
            return jsonify(messages)
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/upload", methods=["POST"])
    def upload_file():
        if "file" not in request.files:
            return jsonify({"error": "未找到文件"}), 400
        file = request.files["file"]
        if file.filename == "":
            return jsonify({"error": "文件名为空"}), 400

        # 支持的文件扩展名
        allowed_extensions = (
            ".txt", ".docx", ".doc", ".pdf", ".xlsx", ".xls",
            ".csv", ".odt", ".ods", ".rtf", ".html", ".htm",
            ".md", ".markdown", ".json", ".xml"
        )
        if not file.filename.lower().endswith(allowed_extensions):
            return jsonify({
                "error": f"不支持的文件格式。支持的格式：{', '.join(allowed_extensions)}"
            }), 400

        try:
            messages = parse_chat_file(file)
            if not messages:
                return jsonify({"error": "文件解析后未找到有效消息"}), 400
            return jsonify(messages)
        except Exception as e:
            return jsonify({"error": f"解析失败：{str(e)}"}), 500

    @app.route("/summarize", methods=["POST"])
    def summarize():
        data = request.get_json()
        messages = data.get("messages", [])
        provider = data.get("provider", "deepseek")
        model = data.get("model", "")
        style = data.get("style", "").strip()
        api_key = data.get("api_key", "").strip()
        keywords_raw = data.get("keywords", "")

        if not api_key:
            saved = get_api_key(provider)
            if saved:
                api_key = saved
            else:
                return jsonify({"error": f"未提供 {provider} 的 API 密钥"}), 400

        style = style if style else None
        if not messages:
            return jsonify({"error": "消息列表不能为空"}), 400

        keywords = [kw.strip() for kw in keywords_raw.split(",") if kw.strip()] if keywords_raw else None

        chat_log = ""
        for msg in messages:
            sender = msg.get("sender", "未知")
            content = msg.get("content", "")
            timestamp = msg.get("time", 0)
            if timestamp:
                time_str = datetime.datetime.fromtimestamp(timestamp).strftime("%m-%d %H:%M")
                chat_log += f"[{time_str}] {sender}: {content}\n"
            else:
                chat_log += f"{sender}: {content}\n"

        try:
            client, default_model = create_client(provider, api_key)
            if not model:
                model = default_model
        except Exception as e:
            return jsonify({"error": str(e)}), 500

        def generate():
            try:
                token_estimate = estimate_token_count(chat_log)
                if token_estimate <= MAX_CHUNK_TOKENS:
                    yield from sse_data("正在总结...")
                    final = summarize_chunk(client, model, chat_log, style, keywords)
                    yield from sse_data(final)
                else:
                    yield from sse_data(f"文本较长（约{token_estimate} tokens），分块处理中...")
                    chunks = split_text_into_chunks(chat_log, MAX_CHUNK_TOKENS)
                    yield from sse_data(f"共 {len(chunks)} 块，正在逐块总结...")
                    summaries = []
                    for i, chunk in enumerate(chunks, 1):
                        yield from sse_data(f"处理第 {i}/{len(chunks)} 块...")
                        summaries.append(summarize_chunk(client, model, chunk, style, keywords))
                    yield from sse_data("正在合并所有摘要...")
                    final = merge_summaries(client, model, summaries, style, keywords)
                    yield from sse_data(final)
                yield from sse_data("[DONE]")
            except Exception as e:
                yield from sse_data(f"[错误] {str(e)}")

        return Response(generate(), mimetype="text/event-stream")

    @app.route("/save_key", methods=["POST"])
    def save_key():
        data = request.get_json()
        provider = data.get("provider", "").strip()
        api_key = data.get("api_key", "").strip()
        if not provider or not api_key:
            return jsonify({"error": "提供商和密钥不能为空"}), 400
        try:
            set_api_key(provider, api_key)
            return jsonify({"success": True})
        except Exception as e:
            return jsonify({"error": str(e)}), 500

    @app.route("/get_key/<provider>")
    def get_key(provider):
        key = get_api_key(provider)
        return jsonify({"key": key})

    # ==================== 一键打包 ====================
    def build_exe():
        ensure_wx_cli()
        print("开始打包...")
        cmd = [
            sys.executable, "-m", "PyInstaller",
            "--onefile",
            "--name", "ChatSummary",
        ]
        cmd.append("--add-data")
        cmd.append(f"templates{os.pathsep}templates")
        if os.path.exists("wx-cli.exe"):
            cmd.append("--add-data")
            cmd.append(f"wx-cli.exe{os.pathsep}.")
        # 隐藏导入多格式解析库
        for mod in ["docx", "doc2txt", "pdfplumber", "openpyxl", "xlrd", "odfdo"]:
            cmd += ["--hidden-import", mod]
        cmd.append("app.py")
        subprocess.check_call(cmd, cwd=BASE_DIR)
        print("打包成功！dist/ChatSummary.exe")

    # ==================== 启动 ====================
    if __name__ == "__main__":
        parser = argparse.ArgumentParser()
        parser.add_argument("--build", action="store_true")
        args = parser.parse_args()
        if args.build:
            build_exe()
            sys.exit(0)

        if not os.path.exists(KEY_FILE):
            generate_key()

        # 释放并初始化 wx-cli（如果存在）
        ensure_wx_cli()
        if os.path.exists(WX_CLI_PATH):
            init_wx_config()
        else:
            print("提示：未找到 wx-cli.exe，自动抓取功能不可用。文件上传和手动粘贴仍可正常使用。")

        if os.environ.get("WERKZEUG_RUN_MAIN") != "true":
            webbrowser.open("http://127.0.0.1:5000")
        app.run(debug=False, host="127.0.0.1", port=5000)

except Exception as e:
    log_crash(e)
    sys.exit(1)
